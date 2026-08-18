#!/usr/bin/env python3
"""按人工分析模板（招标项目分析报告）严格渲染 report-data.yaml 为 docx。

版式严格对照模板：分节标题行、行名、合并单元格与原模板一致。
输入 report-data.yaml 是"凝练层"（由 bid-parse 流程中的模型从
tender-analysis.yaml 逐条凝练产出），本脚本只负责渲染，不改写内容。
"""
import argparse
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

PENDING = "[待补：人工填写]"


def set_cjk(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)


def make_table(doc, ncols):
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    return t


def add_row(t, texts, merge_span=None, bold_cells=()):
    """texts: 每列文本列表；merge_span: (start, end) 列合并区间（含端点）。"""
    row = t.add_row()
    for j, txt in enumerate(texts):
        row.cells[j].text = str(txt) if txt is not None else ""
    if merge_span:
        a, b = merge_span
        row.cells[a].merge(row.cells[b])
    for j in bold_cells:
        for p in row.cells[j].paragraphs:
            for r in p.runs:
                r.bold = True
    return row


def section_row(t, ncols, title):
    row = add_row(t, [title] + [""] * (ncols - 1), merge_span=(0, ncols - 1), bold_cells=(0,))
    return row


def pz(v):
    """空值 → 待补标记；非空原样返回。"""
    return v if (v is not None and str(v).strip() != "") else PENDING


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    ap = argparse.ArgumentParser()
    ap.add_argument("-d", "--data", required=True, help="report-data.yaml 路径")
    ap.add_argument("-o", "--output", required=True, help="输出 docx 路径")
    args = ap.parse_args()

    d = yaml.safe_load(Path(args.data).read_text(encoding="utf-8")) or {}
    info = d.get("info") or {}

    doc = Document()
    set_cjk(doc)
    add_title(doc, "招标项目分析报告")

    # ══ 表一：项目信息 / 招标公告 / 招标文件 / 商务评分（5 列）══
    t1 = make_table(doc, 5)
    add_row(t1, ["项目名称：", info.get("project", "")], merge_span=(1, 4), bold_cells=(0,))
    add_row(t1, ["", "项目来源：", pz(info.get("source")), "渠道：", pz(info.get("channel"))], bold_cells=(1,))
    add_row(t1, ["", "销售经理：", pz(info.get("sales_managers")), "售前经理：" + pz(info.get("presales_manager")), ""], bold_cells=(1,))

    section_row(t1, 5, "招标公告")
    add_row(t1, ["序号", "", "招标文件要求", "招标文件要求", "备注"], bold_cells=(0, 2, 3))
    for no, item, text, note in d.get("announcement") or []:
        add_row(t1, [no, item, text, None, pz(note)], merge_span=(2, 3))

    section_row(t1, 5, "招标文件")
    for item, text, note in d.get("tender_file") or []:
        add_row(t1, ["", item, text, None, pz(note)], merge_span=(2, 3), bold_cells=(1,))

    section_row(t1, 5, "商务评分规则&标准：100分")
    for no, item, text, note in d.get("business_scoring") or []:
        add_row(t1, [no, item, text, None, pz(note)], merge_span=(2, 3))

    # ══ 表二：技术评分 / 价格评分 / 评分办法 / 竞品 / 结论（5 列同栅格）══
    t2 = make_table(doc, 5)
    section_row(t2, 5, "技术评分规则&标准100分")
    for no, item, text, note in d.get("tech_scoring") or []:
        add_row(t2, [no, item, text, None, pz(note)], merge_span=(2, 3))

    section_row(t2, 5, "价格评分规则&标准100分")
    for item, text, note in d.get("price_scoring") or []:
        add_row(t2, ["", item, text, None, pz(note)], merge_span=(2, 3), bold_cells=(1,))

    section_row(t2, 5, "评分办法：综合评估法")
    m = d.get("method") or {}
    add_row(t2, ["", "评分权重", m.get("weights", ""), None, PENDING], merge_span=(2, 3), bold_cells=(1,))
    for label, key in (("商务分析", "biz_analysis"), ("技术分析", "tech_analysis"), ("价格分析", "price_analysis")):
        add_row(t2, ["", label, pz(m.get(key)), None, PENDING], merge_span=(2, 3), bold_cells=(1,))

    section_row(t2, 5, "竞品分析")
    for item, text in d.get("competition") or []:
        add_row(t2, ["", item, pz(text), None, ""], merge_span=(2, 3), bold_cells=(1,))

    section_row(t2, 5, "结论：")
    concl = add_row(t2, [pz(d.get("conclusion")), None, None, None, ""], merge_span=(0, 4))
    for p in concl.cells[0].paragraphs:
        for r in p.runs:
            r.bold = False

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"已生成: {out}")


if __name__ == "__main__":
    main()
